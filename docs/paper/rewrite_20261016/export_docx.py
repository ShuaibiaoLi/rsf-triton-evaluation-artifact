#!/usr/bin/env python3
"""Export the markdown manuscript to a reviewable DOCX with embedded figures."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).resolve().parent
SRC = HERE / "manuscript_cn_rewritten_v1.md"
OUT = HERE / "manuscript_cn_rewritten_v1.docx"

def add_text(doc, text):
    p = doc.add_paragraph()
    add_inline(p, text)
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.35


def add_inline(paragraph, text):
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^]]+\]\([^)]+\))", text):
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2].replace("`", "")); run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1]); run.font.name = "Consolas"
        elif re.fullmatch(r"\[[^]]+\]\([^)]+\)", part):
            label, url = re.fullmatch(r"\[([^]]+)\]\(([^)]+)\)", part).groups()
            paragraph.add_run(f"{label} ({url})")
        elif part:
            paragraph.add_run(part)


def table_cells(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def add_table(doc, lines):
    rows = [table_cells(line) for line in lines]
    width = len(rows[0])
    if len(rows) < 3 or any(len(row) != width for row in rows):
        raise ValueError("invalid Markdown table shape")
    if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        raise ValueError("Markdown table separator missing")
    table = doc.add_table(rows=len(rows) - 1, cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, source_row in enumerate([rows[0], *rows[2:]]):
        for cell, content in zip(table.rows[index].cells, source_row):
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1.1
            add_inline(paragraph, content)
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
                if index == 0:
                    run.bold = True
    header = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header.append(repeat)

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Cm(2.2)
sec.left_margin = sec.right_margin = Cm(2.4)
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
style.font.size = Pt(10.5)

lines = SRC.read_text(encoding="utf-8").splitlines()
index = 0
while index < len(lines):
    raw = lines[index]
    index += 1
    line = raw.strip()
    if not line:
        continue
    if line.startswith("| "):
        block = [line]
        while index < len(lines) and lines[index].strip().startswith("|"):
            block.append(lines[index].strip())
            index += 1
        add_table(doc, block)
        continue
    image = re.fullmatch(r"!\[[^]]*\]\(([^)]+)\)", line)
    if image:
        path = HERE / image.group(1)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Cm(16.0))
        continue
    if line.startswith("# "):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line[2:]); r.bold = True; r.font.size = Pt(16)
        continue
    if line.startswith("## "):
        doc.add_heading(line[3:], level=1); continue
    if line.startswith("### "):
        doc.add_heading(line[4:], level=2); continue
    if line.startswith("!"):
        continue
    if re.fullmatch(r"\d+\. .*", line):
        p = doc.add_paragraph(style="List Number")
        add_inline(p, line.split(". ", 1)[1])
        continue
    add_text(doc, line)

doc.save(OUT)
print(OUT)
